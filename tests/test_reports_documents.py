import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from openpyxl import load_workbook

from apps.documents.models import Document
from apps.documents.search import Criteria, search
from apps.ledger.models import Account, Side, TaxCategory
from apps.ledger.services import LineInput, post_journal
from apps.masters.doc_services import issue_document, save_document
from apps.masters.models import Partner
from apps.reports.selectors import balance_sheet, profit_and_loss, tax_summary
from apps.sales.models import Invoice
from apps.sales.services import post_invoice

D = datetime.date
pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def media(tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path


def acc(name):
    return Account.objects.filter(system_key=name).first() or Account.objects.get(name=name)


def tax(code):
    return TaxCategory.objects.get(code=code)


@pytest.fixture
def books(fiscal_years):
    post_journal(
        D(2025, 4, 1),
        "資本金",
        [
            LineInput(acc("bank"), Side.DEBIT, Decimal(3_000_000)),
            LineInput(acc("capital"), Side.CREDIT, Decimal(3_000_000)),
        ],
    )
    post_journal(
        D(2025, 6, 1),
        "前期売上",
        [
            LineInput(acc("bank"), Side.DEBIT, Decimal(1_100_000)),
            LineInput(acc("sales"), Side.CREDIT, Decimal(1_100_000), tax("S10")),
        ],
    )
    post_journal(
        D(2026, 5, 1),
        "売上",
        [
            LineInput(acc("accounts_receivable"), Side.DEBIT, Decimal(2_200_000)),
            LineInput(acc("sales"), Side.CREDIT, Decimal(2_200_000), tax("S10")),
        ],
    )
    post_journal(
        D(2026, 5, 10),
        "外注",
        [
            LineInput(acc("外注費"), Side.DEBIT, Decimal(550_000), tax("P10")),
            LineInput(acc("accounts_payable"), Side.CREDIT, Decimal(550_000)),
        ],
    )
    post_journal(
        D(2026, 5, 20),
        "家賃",
        [
            LineInput(acc("地代家賃"), Side.DEBIT, Decimal(220_000), tax("P10")),
            LineInput(acc("bank"), Side.CREDIT, Decimal(220_000)),
        ],
    )
    post_journal(
        D(2026, 5, 25),
        "利息",
        [
            LineInput(acc("bank"), Side.DEBIT, Decimal(100)),
            LineInput(acc("受取利息"), Side.CREDIT, Decimal(100), tax("SNT")),
        ],
    )


def test_profit_and_loss(books):
    lines = {ln.label: ln for ln in profit_and_loss(D(2026, 4, 1), D(2027, 3, 31), compare=True)}
    assert lines["売上高 合計"].amount == 2_000_000
    assert lines["売上高 合計"].compare == 1_000_000
    assert lines["売上総利益"].amount == 2_000_000 - 500_000
    assert lines["営業利益"].amount == 1_500_000 - 200_000
    assert lines["経常利益"].amount == 1_300_100
    assert lines["当期純利益"].amount == 1_300_100


def test_balance_sheet_balances(books):
    bs = balance_sheet(D(2026, 6, 30))
    assert bs.total_assets == bs.total_liabilities_equity
    labels = {ln.label: ln.amount for ln in bs.liabilities}
    assert labels["繰越利益剰余金"] == 1_000_000  # 前期利益
    assert labels["当期純利益"] == 1_300_100
    assert labels["仮受消費税"] == 300_000  # 前期分100,000（納付の仕訳なし）＋当期分


def test_tax_summary(books):
    s = tax_summary(D(2026, 4, 1), D(2027, 3, 31))
    assert s.taxable_sales == 2_000_000 and s.sales_tax == 200_000
    assert s.exempt_sales == 100
    assert s.purchase_tax == 70_000 and s.net_tax == 130_000
    gross, standard, rate = s.sales_tax_by_rate[10]
    assert gross == 2_200_000 and standard == 2_000_000
    assert s.taxable_sales_ratio == Decimal(2_000_000) / Decimal(2_000_100)


def test_documents_storage_and_search(admin_client, fiscal_years):
    partner = Partner.objects.create(code="C1", name="株式会社検索テスト")
    resp = admin_client.post(
        "/documents/new/",
        {
            "kind": "contract",
            "source": "electronic",
            "transaction_date": "2026-05-01",
            "amount": "330000",
            "partner": partner.pk,
            "title": "業務委託契約書",
            "file": SimpleUploadedFile("c.pdf", b"%PDF contract"),
        },
    )
    assert resp.status_code == 302, resp.content.decode()[:2000]
    doc = Document.objects.get()
    assert len(doc.file_sha256) == 64
    with pytest.raises(ValueError):
        doc.delete()
    doc.file = SimpleUploadedFile("other.pdf", b"x")
    with pytest.raises(ValueError):
        doc.save()
    inv = issue_document(
        save_document(
            Invoice(partner=partner, issue_date=D(2026, 5, 31)),
            [{"description": "x", "quantity": Decimal(1), "unit_price": Decimal(100000), "tax_rate": "10"}],
        ),
        "pdf/invoice.html",
        on_issue=post_invoice,
    )
    hits = search(Criteria(partner="検索テスト"))
    assert {h.kind for h in hits} == {"契約書", "請求書（発行）"}
    assert [h.number for h in search(Criteria(amount_min=Decimal(100000), amount_max=Decimal(200000)))] == [
        inv.number
    ]
    assert search(Criteria(date_from=D(2026, 6, 1))) == []
    resp = admin_client.get("/documents/", {"partner": "検索テスト", "go": "1"})
    assert "業務委託契約書" in resp.content.decode()
    admin_client.post(f"/documents/{doc.pk}/", {"reason": "重複登録"})
    doc.refresh_from_db()
    assert doc.is_void
    assert admin_client.get(f"/documents/{doc.pk}/file/").status_code == 200
    call_command("verify_files")
    # 改ざんを検知する
    with open(doc.file.path, "ab") as f:
        f.write(b"tampered")
    with pytest.raises(SystemExit):
        call_command("verify_files")


def test_report_views_and_excel(admin_client, books):
    for url in [
        "/",
        "/reports/pl/?compare=1",
        "/reports/bs/",
        "/reports/tax/",
        "/reports/pl/?format=csv",
        "/documents/",
        "/ledger/trial-balance/?format=csv",
    ]:
        assert admin_client.get(url).status_code == 200, url
    resp = admin_client.get("/reports/bs/?as_of=2026-06-30&format=xlsx")
    ws = load_workbook(BytesIO(resp.content)).active
    values = {row[0].strip("　"): row[1] for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}
    assert values["資産の部合計"] == values["負債・純資産合計"]
    resp = admin_client.get("/ledger/trial-balance/?date_from=2026-04-01&date_to=2026-06-30&format=xlsx")
    assert resp["Content-Type"].startswith("application/vnd.openxmlformats")

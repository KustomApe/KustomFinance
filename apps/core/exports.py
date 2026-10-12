import csv
from urllib.parse import quote

from django.http import HttpResponse


def csv_response(filename: str, header: list[str], rows) -> HttpResponse:
    """Excelで文字化けしないよう BOM 付き UTF-8 で CSV を返す。"""
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    response.write("﻿")
    writer = csv.writer(response)
    writer.writerow(header)
    writer.writerows(rows)
    return response


def xlsx_response(filename: str, header: list[str], rows, title: str = "Sheet1") -> HttpResponse:
    """Excel（.xlsx）で返す。数値はそのまま数値セルにする。"""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEF2F7")
    for row in rows:
        ws.append([float(v) if hasattr(v, "is_finite") else v for v in row])
    for column in ws.columns:
        width = max(len(str(c.value or "")) for c in column)
        ws.column_dimensions[column[0].column_letter].width = min(max(width * 1.8, 8), 60)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0"
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return response


def table_response(fmt: str, filename: str, header: list[str], rows):
    """format=csv / xlsx に応じて出力する。"""
    rows = list(rows)
    if fmt == "xlsx":
        return xlsx_response(f"{filename}.xlsx", header, rows, title=filename)
    return csv_response(f"{filename}.csv", header, rows)
